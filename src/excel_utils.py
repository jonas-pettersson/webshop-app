"""Excel template manipulation utilities.

These functions handle the sophisticated Excel template system with named ranges,
dynamic row insertion, and format preservation.
"""

from copy import copy
from openpyxl import Workbook


def get_first_cell_for_named_range(workbook, sheet, range_name):
    """Extract the first cell coordinate from a workbook-level named range.

    Args:
        workbook: openpyxl Workbook object
        sheet: openpyxl Worksheet object
        range_name: Name of the defined range in the workbook

    Returns:
        str: Cell coordinate (e.g., 'A1') or empty string if not found
    """
    cell_name_list = []
    first_cell_tuple = ()

    if range_name in workbook.defined_names:
        cell_name_list = [
            (s, c) for s, c in workbook.defined_names[range_name].destinations
            if (s == sheet.title)
        ]

    if cell_name_list:
        first_cell_tuple = cell_name_list[0]
        return first_cell_tuple[1]

    return ''


def fill_cell_for_attribute(range_name, workbook, template, sheet, row_offset, value):
    """Fill a cell using named range lookup with row offset.

    This function locates a named range in the template, applies a row offset,
    and fills the corresponding cell in the target sheet with the provided value.

    Args:
        range_name: Name of the defined range in the workbook
        workbook: openpyxl Workbook object
        template: Template worksheet (source)
        sheet: Target worksheet (destination)
        row_offset: Number of rows to offset from the named range position
        value: Value to write to the cell
    """
    template_cell_number = get_first_cell_for_named_range(workbook, template, range_name)
    if not template_cell_number:
        return

    target_cell_number = template[template_cell_number].offset(row=row_offset).coordinate
    sheet[target_cell_number] = value


def get_merged_cells_or_none(sheet, row, column):
    """Check if a cell is part of a merged cell range.

    Args:
        sheet: openpyxl Worksheet object
        row: Row number (1-indexed)
        column: Column number (1-indexed)

    Returns:
        MergedCellRange object if cell is merged, None otherwise
    """
    cell = sheet.cell(row, column)

    for merged_cell in sheet.merged_cells.ranges:
        if cell.coordinate in merged_cell:
            return merged_cell

    return None


def insert_row_and_copy_format(template, sheet, source_row, target_row):
    """Insert a row in the sheet and copy all formatting from the template.

    This preserves cell styles, row height, and merged cell ranges from the template.
    Critical for maintaining Excel report appearance.

    Args:
        template: Template worksheet (source of formatting)
        sheet: Target worksheet (where row is inserted)
        source_row: Row number in template to copy formatting from
        target_row: Row number in sheet where new row should be inserted
    """
    sheet.insert_rows(idx=target_row)
    sheet.row_dimensions[target_row].height = template.row_dimensions[source_row].height

    # Copy cell styles and handle merged cells
    for col in range(1, 11):  # Assuming max 10 columns
        sheet.cell(row=target_row, column=col)._style = template.cell(row=source_row, column=col)._style

        merged_cells = get_merged_cells_or_none(template, source_row, col)
        if merged_cells is not None and col == merged_cells.bounds[0]:
            # Only process the first column of a merged range
            merged_cells_copy = copy(merged_cells)
            if target_row > source_row:
                merged_cells_copy.shift(row_shift=target_row - source_row)
                sheet.merge_cells(range_string=(merged_cells_copy.coord))


def align_cell_width_for(sheet):
    """Auto-adjust column widths based on cell content.

    Iterates through all cells and sets column width to accommodate the longest
    content in each column.

    Args:
        sheet: openpyxl Worksheet object
    """
    dims = {}

    for row in sheet.rows:
        for cell in row:
            if cell.value:
                dims[cell.column_letter] = max(
                    (dims.get(cell.column_letter, 0), len(str(cell.value)))
                )

    for col, value in dims.items():
        sheet.column_dimensions[col].width = value * 1.23
