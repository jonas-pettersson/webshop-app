"""Data validation for WebShop Report system.

Validates relationships between orders, timeslots, products, and packing categories.
Generates comprehensive validation report (diffs.xlsx).
"""

from pathlib import Path
from typing import Dict

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .excel_utils import align_cell_width_for


class ValidationResult:
    """Container for validation results."""

    def __init__(self):
        self.orders_no_timeslot = pd.DataFrame()
        self.timeslots_no_order = pd.DataFrame()
        self.duplicate_bookings = pd.DataFrame()
        self.missing_products = pd.DataFrame()
        self.missing_categories = pd.DataFrame()

    def has_errors(self) -> bool:
        """Check if any validation errors were found."""
        return (
            len(self.orders_no_timeslot) > 0 or
            len(self.timeslots_no_order) > 0 or
            len(self.duplicate_bookings) > 0 or
            len(self.missing_products) > 0 or
            len(self.missing_categories) > 0
        )

    def get_summary(self) -> Dict[str, int]:
        """Get count of each validation issue type."""
        return {
            'orders_no_timeslot': len(self.orders_no_timeslot),
            'timeslots_no_order': len(self.timeslots_no_order),
            'duplicate_bookings': len(self.duplicate_bookings),
            'missing_products': len(self.missing_products),
            'missing_categories': len(self.missing_categories)
        }


def validate_data(
    orders_df: pd.DataFrame,
    timeslots_df: pd.DataFrame,
    products_df: pd.DataFrame,
    packing_df: pd.DataFrame
) -> ValidationResult:
    """Validate relationships between all datasets.

    Checks for:
    - Orders without time bookings
    - Time bookings without orders
    - Duplicate time bookings (multiple per email)
    - Products in orders but not in product catalog
    - Packing categories in orders but not defined

    Args:
        orders_df: Orders DataFrame
        timeslots_df: Timeslots DataFrame
        products_df: Products DataFrame
        packing_df: Packing categories DataFrame

    Returns:
        ValidationResult object containing all validation issues
    """
    result = ValidationResult()

    # 1. Orders without time bookings
    order_emails = orders_df['Email'].unique()
    timeslot_emails = timeslots_df['Email'].unique()
    diff_emails = np.setdiff1d(order_emails, timeslot_emails)

    result.orders_no_timeslot = orders_df[
        orders_df['Email'].isin(diff_emails)
    ].drop_duplicates(subset=['OrderNumber'])[[
        'OrderNumber', 'FullName', 'Email',
        'OrderTotalAmount', 'OrderDate', 'CustomerNote', 'PhoneOrder'
    ]]

    # 2. Time bookings without orders
    diff_emails = np.setdiff1d(timeslot_emails, order_emails)
    result.timeslots_no_order = timeslots_df[
        timeslots_df['Email'].isin(diff_emails)
    ][[
        'Email', 'PickUpDate', 'PickUpTime', 'FullNameTimeslot', 'Telephone'
    ]]

    # 3. Duplicate time bookings
    email_counts = timeslots_df['Email'].value_counts()
    duplicate_emails = email_counts.index[email_counts.gt(1)]
    result.duplicate_bookings = timeslots_df[
        timeslots_df['Email'].isin(duplicate_emails)
    ].sort_values(by=['Email', 'PickUpDate', 'PickUpTime'])[[
        'Email', 'PickUpDate', 'PickUpTime', 'FullNameTimeslot', 'Telephone'
    ]]

    # 4. Products in orders but missing from product catalog
    order_skus = orders_df['ArticleNumber'].unique()
    product_skus = products_df['ArticleNumber'].unique()
    sku_diff = np.setdiff1d(order_skus, product_skus)

    result.missing_products = orders_df[
        orders_df['ArticleNumber'].isin(sku_diff)
    ][['ArticleNumber', 'ProductName']].drop_duplicates(
        subset=['ArticleNumber']
    ).sort_values(by='ArticleNumber')

    # 5. Packing categories in orders but not defined
    order_categories = orders_df['PackingCategoryKey'].unique()
    packing_categories = packing_df['PackingCategoryKey'].unique()
    category_diff = np.setdiff1d(order_categories, packing_categories)

    result.missing_categories = orders_df[
        orders_df['PackingCategoryKey'].isin(category_diff)
    ][[
        'PackingCategory', 'OrderNumber', 'ArticleNumber',
        'ProductName', 'FullName', 'Email'
    ]].sort_values(by='PackingCategory')

    return result


def generate_diffs_report(validation_result: ValidationResult, output_path: Path):
    """Generate Excel validation report with multiple sheets.

    Creates diffs.xlsx with 5 sheets showing validation issues.

    Args:
        validation_result: ValidationResult object from validate_data()
        output_path: Path where diffs.xlsx should be written
    """
    # Remove existing file
    if output_path.exists():
        output_path.unlink()

    # Sheet 1: Orders without time bookings
    validation_result.orders_no_timeslot.to_excel(
        output_path,
        sheet_name='Beställning utan tidsbokning'
    )

    workbook = load_workbook(filename=output_path)
    sheet = workbook['Beställning utan tidsbokning']
    align_cell_width_for(sheet)
    workbook.save(filename=output_path)

    # Sheet 2: Time bookings without orders
    with pd.ExcelWriter(output_path, mode='a') as writer:
        validation_result.timeslots_no_order.to_excel(
            writer,
            sheet_name='Tidsbokning utan beställning'
        )

    workbook = load_workbook(filename=output_path)
    sheet = workbook['Tidsbokning utan beställning']
    align_cell_width_for(sheet)
    workbook.save(filename=output_path)

    # Sheet 3: Duplicate time bookings
    with pd.ExcelWriter(output_path, mode='a') as writer:
        validation_result.duplicate_bookings.to_excel(
            writer,
            sheet_name='Dubbla tidsbokningar'
        )

    workbook = load_workbook(filename=output_path)
    sheet = workbook['Dubbla tidsbokningar']
    align_cell_width_for(sheet)
    workbook.save(filename=output_path)

    # Sheet 4: Missing products
    with pd.ExcelWriter(output_path, mode='a') as writer:
        validation_result.missing_products.to_excel(
            writer,
            sheet_name='Saknas i produktlista'
        )

    workbook = load_workbook(filename=output_path)
    sheet = workbook['Saknas i produktlista']
    align_cell_width_for(sheet)
    workbook.save(filename=output_path)

    # Sheet 5: Missing packing categories
    with pd.ExcelWriter(output_path, mode='a') as writer:
        validation_result.missing_categories.to_excel(
            writer,
            sheet_name='Packningskategori saknas'
        )

    workbook = load_workbook(filename=output_path)
    sheet = workbook['Packningskategori saknas']
    align_cell_width_for(sheet)
    workbook.save(filename=output_path)


def remove_duplicate_bookings(timeslots_df: pd.DataFrame) -> pd.DataFrame:
    """Remove duplicate time bookings, keeping only the first booking per email.

    This should be called AFTER validation report is generated, so duplicates
    are reported but then removed for processing.

    Args:
        timeslots_df: Timeslots DataFrame (will not be modified)

    Returns:
        New DataFrame with duplicates removed
    """
    return timeslots_df.drop_duplicates(subset=['Email'], keep='first')
