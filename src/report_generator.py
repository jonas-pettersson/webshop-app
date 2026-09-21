"""Report generation for WebShop Report system.

Generates Excel reports using template-based approach:
- Packlista (packing lists per order)
- Hämtningslista (pickup schedules by date)
- Semlelista (special product list for pask market)
- Masterlista (master order list)
"""

from datetime import datetime
from pathlib import Path
import shutil

import pandas as pd
from openpyxl import load_workbook

from . import config
from .excel_utils import (
    get_first_cell_for_named_range,
    fill_cell_for_attribute,
    insert_row_and_copy_format
)


# Weekday names in Swedish
WEEKDAY_DICT = {
    0: 'Mån',
    1: 'Tis',
    2: 'Ons',
    3: 'Tor',
    4: 'Fre',
    5: 'Lör',
    6: 'Sön'
}


def generate_packlista(
    all_data_df: pd.DataFrame,
    template_dir: Path,
    output_path: Path,
    market_type: str
):
    """Generate packing lists (one sheet per order).

    Creates packlista.xlsx with one worksheet per order, grouped by packing area.

    Args:
        all_data_df: Merged and sorted DataFrame
        template_dir: Path to templates directory
        output_path: Path where packlista.xlsx should be written
        market_type: One of config.MARKET_TYPES
    """
    # Copy template to output
    template_path = config.get_template_path(
        template_dir,
        config.TEMPLATE_FILENAME_PACKLISTA,
        market_type
    )

    if output_path.exists():
        output_path.unlink()

    shutil.copyfile(template_path, output_path)

    # Load workbook and template sheet
    workbook = load_workbook(filename=output_path)
    template_sheet = workbook['template']

    # Get header row position
    article_cell = get_first_cell_for_named_range(workbook, template_sheet, 'ArticleNumber')
    header_row = template_sheet[article_cell].row

    # Group by order number (already sorted by pickup date/time)
    df_grouped_by_order = all_data_df.groupby(
        by=['PickUpDate', 'PickUpTime', 'OrderNumber']
    )

    for index, group in df_grouped_by_order:
        # Create new sheet for this order
        new_sheet = workbook.copy_worksheet(template_sheet)
        order_number = group.iloc[0]['OrderNumber']
        new_sheet.title = 'PO_' + str(order_number)

        # Fill header information
        for range_name in ['FullName', 'OrderTotalAmount', 'CustomerNote',
                          'OrderNumber', 'Telephone', 'PickUpDate', 'PickUpTime']:
            field_value = group.iloc[0][range_name]
            fill_cell_for_attribute(range_name, workbook, template_sheet, new_sheet, 0, field_value)

        # Delete template rows
        new_sheet.delete_rows(idx=header_row, amount=2)

        # Group by pickup area (sort=False preserves template order)
        df_grouped_by_area = group.groupby('PickUpArea', sort=False)

        i = 0  # Row counter
        d = 0  # Deleted row counter (for row numbering)

        for area, sub_group in df_grouped_by_area:
            # Add rows for each item in this area
            for sub_index, row in sub_group.iterrows():
                insert_row_and_copy_format(template_sheet, new_sheet, header_row, header_row + i)

                # Fill item data
                for range_name in ['ArticleNumber', 'Quantity', 'ProductName', 'ItemCost', 'PickUpArea']:
                    field_value = row[range_name]
                    fill_cell_for_attribute(range_name, workbook, template_sheet, new_sheet, i, field_value)

                i += 1
                fill_cell_for_attribute('RowNumber', workbook, template_sheet, new_sheet, i - 1, i - d)

            # Add separator row after each area
            insert_row_and_copy_format(template_sheet, new_sheet, header_row + 1, header_row + i)
            i += 1
            d += 1

        # Copy footer formatting
        for j in range(0, 10):
            offset = header_row + i
            new_sheet.row_dimensions[offset + j].height = template_sheet.row_dimensions[header_row + 2 + j].height

        # Fill total number of rows
        fill_cell_for_attribute('NumberOfRows', workbook, template_sheet, new_sheet, 0, str(i - d))

    # Remove template sheet and save
    workbook.remove(template_sheet)
    workbook.save(filename=output_path)


def generate_hamtningslista(
    all_data_df: pd.DataFrame,
    schedule_df: pd.DataFrame,
    template_dir: Path,
    output_path: Path
):
    """Generate pickup lists (one sheet per pickup date).

    Creates hämtningslista.xlsx with one worksheet per date, organized by time slots.

    Args:
        all_data_df: Merged and sorted DataFrame
        schedule_df: Schedule DataFrame with authorized pickup times
        template_dir: Path to templates directory
        output_path: Path where hämtningslista.xlsx should be written
    """
    # Copy template to output
    template_path = template_dir / config.TEMPLATE_FILENAME_HAMTLISTA

    if output_path.exists():
        output_path.unlink()

    shutil.copyfile(template_path, output_path)

    # Load workbook and template sheet
    workbook = load_workbook(filename=output_path)
    template_sheet = workbook['template']

    # Get header row position
    pickup_time_cell = get_first_cell_for_named_range(workbook, template_sheet, 'PickUpTime')
    header_row = template_sheet[pickup_time_cell].row

    # Group by pickup date (one sheet per date)
    df_grouped_by_date = all_data_df.drop_duplicates(subset=['OrderNumber']).groupby('PickUpDate')

    for pickup_date, group in df_grouped_by_date:
        # Create new sheet for this date
        new_sheet = workbook.copy_worksheet(template_sheet)
        new_sheet.title = 'Hämtas ' + str(pickup_date.date())

        # Fill date information
        fill_cell_for_attribute('Weekday', workbook, template_sheet, new_sheet, 0, WEEKDAY_DICT[pickup_date.weekday()])
        fill_cell_for_attribute('PickUpDate', workbook, template_sheet, new_sheet, 0, pickup_date.date())
        fill_cell_for_attribute('CreationDateTime', workbook, template_sheet, new_sheet, 0, str(datetime.now()))

        # Delete template rows
        new_sheet.delete_rows(idx=header_row, amount=2)

        # Sort by full name
        group = group.sort_values(by=['FullName'])
        df_grouped_by_time = group.groupby('PickUpTime')

        # Get scheduled times for this date
        scheduled_times = schedule_df[schedule_df['Date'] == pickup_date.date()]['Tid']

        i = 0  # Row counter
        d = 0  # Deleted row counter

        for time in scheduled_times:
            # Add time slot header row
            insert_row_and_copy_format(template_sheet, new_sheet, header_row, header_row + i)
            fill_cell_for_attribute('PickUpTime', workbook, template_sheet, new_sheet, i, time)
            i += 1
            d += 1

            # Add orders for this time slot if any exist
            if time in df_grouped_by_time.indices:
                sub_group = df_grouped_by_time.get_group(time)

                for index, row in sub_group.iterrows():
                    insert_row_and_copy_format(template_sheet, new_sheet, header_row + 1, header_row + i)

                    # Fill order data
                    for range_name in ['OrderNumber', 'FullName', 'Telephone', 'OrderTotalAmount', 'CustomerNote']:
                        field_value = row[range_name]
                        fill_cell_for_attribute(range_name, workbook, template_sheet, new_sheet, i - 1, field_value)

                    i += 1
                    fill_cell_for_attribute('RowNumber', workbook, template_sheet, new_sheet, i - 2, i - d)
            else:
                # Empty time slot - add placeholder row
                insert_row_and_copy_format(template_sheet, new_sheet, header_row + 1, header_row + i)
                i += 1
                fill_cell_for_attribute('RowNumber', workbook, template_sheet, new_sheet, i - 2, i - d)

    # Remove template sheet and save
    workbook.remove(template_sheet)
    workbook.save(filename=output_path)


def generate_semlelista(
    all_data_df: pd.DataFrame,
    template_dir: Path,
    output_path: Path
):
    """Generate semla list (special product list for pask market only).

    Creates semlelista.xlsx with orders for semla items grouped by date and hour.

    Args:
        all_data_df: Merged and sorted DataFrame
        template_dir: Path to templates directory
        output_path: Path where semlelista.xlsx should be written
    """
    # Copy template to output
    template_path = template_dir / config.TEMPLATE_FILENAME_SEMLELISTA

    if output_path.exists():
        output_path.unlink()

    shutil.copyfile(template_path, output_path)

    # Load workbook and template sheet
    workbook = load_workbook(filename=output_path)
    template_sheet = workbook['template']

    # Filter for semla orders only
    df_semla = all_data_df[all_data_df['ArticleNumber'] == config.ARTICLE_NUMBER_SEMLA]
    df_semla = df_semla[[
        'ArticleNumber', 'ProductName', 'Quantity', 'OrderNumber',
        'FullName', 'Email', 'PickUpDate', 'PickUpTime'
    ]]

    # Create pickup datetime and extract hour
    df_semla = df_semla.copy()
    df_semla['PickUpDateTime'] = [
        datetime.combine(d, t) for d, t in zip(df_semla['PickUpDate'], df_semla['PickUpTime'])
    ]

    if len(df_semla) > 0:
        df_semla['PickUpHour'] = df_semla['PickUpDateTime'].dt.hour

    # Get header row position
    order_number_cell = get_first_cell_for_named_range(workbook, template_sheet, 'OrderNumber')
    header_row = template_sheet[order_number_cell].row

    # Create single sheet
    new_sheet = workbook.copy_worksheet(template_sheet)
    new_sheet.title = 'Semla-Lista'

    # Fill header information
    for range_name in ['ArticleNumber', 'ProductName']:
        field_value = ''
        if len(df_semla) > 0:
            field_value = df_semla.iloc[0][range_name]
        fill_cell_for_attribute(range_name, workbook, template_sheet, new_sheet, 0, field_value)

    # Delete template rows
    new_sheet.delete_rows(idx=header_row, amount=2)

    i = 0

    # Group by pickup date
    df_grouped_by_date = df_semla.groupby('PickUpDate')

    for day, group in df_grouped_by_date:
        # Group by hour
        df_grouped_by_hour = group.groupby('PickUpHour')

        for hour, sub_group in df_grouped_by_hour:
            # Add individual orders
            for sub_index, row in sub_group.iterrows():
                insert_row_and_copy_format(template_sheet, new_sheet, header_row, header_row + i)

                for range_name in ['Quantity', 'OrderNumber', 'FullName', 'Email', 'PickUpDate', 'PickUpTime']:
                    field_value = row[range_name]
                    fill_cell_for_attribute(range_name, workbook, template_sheet, new_sheet, i, field_value)

                i += 1

            # Add summary row for this hour
            insert_row_and_copy_format(template_sheet, new_sheet, header_row + 1, header_row + i)
            fill_cell_for_attribute('Day', workbook, template_sheet, new_sheet, i - 1, day.strftime('%d.%m'))

            time_interval_string = str(hour) + ':00 - ' + str(hour + 1) + ':00'
            fill_cell_for_attribute('Hour', workbook, template_sheet, new_sheet, i - 1, time_interval_string)
            fill_cell_for_attribute('NumberOfItemsName', workbook, template_sheet, new_sheet, i - 1, 'Antal:')

            number_of_items = sub_group['Quantity'].sum()
            fill_cell_for_attribute('NumberOfItems', workbook, template_sheet, new_sheet, i - 1, number_of_items)

            i += 2

        # Copy footer formatting
        for j in range(0, 10):
            offset = header_row + i
            new_sheet.row_dimensions[offset + j].height = template_sheet.row_dimensions[header_row + 2 + j].height

    # Remove template sheet and save
    workbook.remove(template_sheet)
    workbook.save(filename=output_path)


def generate_masterlista(
    all_data_df: pd.DataFrame,
    output_path: Path
):
    """Generate master list (simple consolidated order list).

    Creates masterlista.xlsx with one row per order.

    Args:
        all_data_df: Merged and sorted DataFrame
        output_path: Path where masterlista.xlsx should be written
    """
    df_master = all_data_df[[
        'OrderNumber', 'FullName', 'Email', 'OrderTotalAmount', 'OrderDate',
        'Telephone', 'CustomerNote', 'Info', 'PickUpDate', 'PickUpTime'
    ]].copy()

    df_master.columns = [
        'Order Number', 'Full Name (Billing)', 'Email (Billing)',
        'Order Total Amount', 'Order Date', 'Telephone', 'Customer Note',
        'Info ok', 'Hämtningsdag', 'Hämtning Tid'
    ]

    # Remove duplicates
    df_master = df_master.drop_duplicates()

    df_master.to_excel(output_path)
